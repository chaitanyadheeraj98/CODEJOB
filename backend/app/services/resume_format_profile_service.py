"""Read an employer's sample resume and turn its layout into a profile.

The split here matters. Margins, fonts and column widths are *measured* out of
the sample with python-docx - they are recorded in the file, so asking a model
to guess them would be inventing numbers that are sitting right there. What the
model is asked for is the one thing the text carries and the file does not: which
of the headings are section headings that want a rule drawn above them. That is a
short list of strings, validated against the document before it is stored, so a
bad answer degrades to the defaults instead of corrupting a layout.

Nothing here generates or runs code. The earlier version of this workflow was a
hand-written Google Apps Script per employer; a profile is the same information
as data, which is why it can be produced automatically and safely.
"""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path

from docx import Document
from pydantic import BaseModel, Field, ValidationError

from app.ai.chat.history import message_text
from app.ai.chat.llm import build_chat_llm
from app.parsing.document_extraction import extract_document_text
from app.services.resume_render_service import ResumeFormatSpec

logger = logging.getLogger(__name__)

_SECTIONS_PROMPT = """List the section headings of this resume as JSON only, with this exact shape:
{{"sections":["Summary","Skills","Experiences"]}}

A section heading is a short line that names a block of the resume, such as
Summary, Skills, Certifications, Experiences, Education Details, or Projects.
Copy each heading exactly as it is written in the resume. Do not translate,
re-order, re-word or invent headings. Do not include the candidate's name,
their contact line, job titles, company names, or dates.

Resume:
{text}
"""

# A heading longer than this is a sentence that happens to sit on its own line.
_MAX_HEADING_WORDS = 5
_MAX_SECTIONS = 12


class _SectionList(BaseModel):
    sections: list[str] = Field(default_factory=list)


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip().strip("#*_:")).strip()


def _emu_to_inches(value, fallback: float) -> float:
    """Word lengths are EMU-typed; a section that never set one reads back None."""
    try:
        return round(float(value.inches), 3)
    except (AttributeError, TypeError, ValueError):
        return fallback


def measure_docx(path: Path, spec: ResumeFormatSpec) -> ResumeFormatSpec:
    """Overwrite `spec`'s geometry with what the sample .docx actually contains.

    Anything the sample does not state keeps the value it already had, so a
    sparse document narrows the profile rather than resetting it.
    """
    document = Document(str(path))
    updates: dict[str, object] = {}

    section = document.sections[0]
    width = _emu_to_inches(section.page_width, spec.page_width_inches)
    height = _emu_to_inches(section.page_height, spec.page_height_inches)
    updates["page_size"] = "A4" if abs(width - 210 / 25.4) < 0.05 and abs(height - 297 / 25.4) < 0.05 else "LETTER"
    updates["margin_left_inches"] = _emu_to_inches(section.left_margin, spec.margin_left_inches)
    updates["margin_right_inches"] = _emu_to_inches(section.right_margin, spec.margin_right_inches)
    updates["margin_top_inches"] = _emu_to_inches(section.top_margin, spec.margin_top_inches)
    updates["margin_bottom_inches"] = _emu_to_inches(section.bottom_margin, spec.margin_bottom_inches)

    normal = document.styles["Normal"].font
    if normal.name:
        updates["font_family"] = normal.name
    if normal.size is not None:
        updates["body_font_size"] = round(float(normal.size.pt), 1)
        updates["heading_font_size"] = round(float(normal.size.pt), 1)

    # The name is the largest run in the document's first few paragraphs, which
    # is a more reliable signal than paragraph order in a converted file.
    sizes = [
        float(run.font.size.pt)
        for paragraph in document.paragraphs[:6]
        for run in paragraph.runs
        if run.font.size is not None
    ]
    if sizes:
        updates["name_font_size"] = round(max(sizes), 1)

    if document.tables:
        first_column = document.tables[0].columns[0]
        divider = _emu_to_inches(first_column.width, spec.skills_divider_inches)
        usable = (210 / 25.4 if updates["page_size"] == "A4" else 8.5) - float(updates["margin_left_inches"]) - float(updates["margin_right_inches"])
        # A width wider than the page means the sample never set one and Word
        # reported its own default; the current divider is the better answer.
        if 0 < divider < usable:
            updates["skills_divider_inches"] = divider

    return spec.model_copy(update=updates)


def _headings_in(text: str) -> set[str]:
    """Every line short enough to be a heading, lower-cased for comparison.

    The model's answer is checked against this. A heading it returns that is not
    actually in the document is a hallucination, and drawing a rule above a
    heading that does not exist would silently do nothing anyway.
    """
    found = set()
    for line in text.replace("\r\n", "\n").split("\n"):
        cleaned = _clean(line)
        if cleaned and len(cleaned.split()) <= _MAX_HEADING_WORDS:
            found.add(cleaned.lower())
    return found


def _validate_sections(raw: str, document_text: str) -> list[str]:
    text = (raw or "").strip()
    fenced = re.fullmatch(r"```(?:json)?\s*(.*?)\s*```", text, flags=re.DOTALL | re.IGNORECASE)
    if fenced:
        text = fenced.group(1).strip()
    try:
        parsed = _SectionList.model_validate(json.loads(text))
    except (ValidationError, ValueError) as exc:
        logger.warning("Resume section JSON was invalid: %s", exc)
        return []

    present = _headings_in(document_text)
    kept: list[str] = []
    for candidate in parsed.sections:
        cleaned = _clean(candidate)
        if not cleaned or len(cleaned.split()) > _MAX_HEADING_WORDS:
            continue
        if cleaned.lower() not in present:
            logger.info("Ignoring section %r: it is not a line in the sample", cleaned)
            continue
        if cleaned.lower() not in {item.lower() for item in kept}:
            kept.append(cleaned)
    return kept[:_MAX_SECTIONS]


def read_sections(document_text: str) -> list[str]:
    """Ask the model which headings are sections. Returns [] if it cannot say."""
    if not document_text.strip():
        return []
    try:
        llm = build_chat_llm()
    except Exception as exc:
        logger.warning("Resume format model unavailable: %s", exc)
        return []
    try:
        response = llm.invoke(_SECTIONS_PROMPT.format(text=document_text[:12000]))
    except Exception as exc:
        logger.warning("Resume section detection skipped: %s", exc)
        return []
    return _validate_sections(message_text(response.content), document_text)


def spec_from_sample(path: Path, file_name: str) -> ResumeFormatSpec:
    """Build a profile from an employer's sample resume.

    Every step is optional and falls back to the default layout, because a
    profile that is half-measured is still more useful than a failed upload.
    """
    spec = ResumeFormatSpec()

    if path.suffix.lower() == ".docx":
        try:
            spec = measure_docx(path, spec)
        except Exception as exc:
            logger.warning("Could not measure %s, using default geometry: %s", file_name, exc)

    try:
        extracted = extract_document_text(str(path), file_name, max_chars=12000, label="Sample resume")
    except Exception as exc:
        logger.warning("Could not read %s: %s", file_name, exc)
        return spec

    sections = read_sections(extracted.markdown_text)
    if sections:
        spec = spec.model_copy(update={"rule_before_sections": sections})
    return spec
