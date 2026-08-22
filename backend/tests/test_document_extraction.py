import base64
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from unstructured.documents.elements import (
    ElementMetadata,
    Footer,
    Header,
    ListItem,
    NarrativeText,
    Table,
    Title,
)

from app.ai.resume_context import extract_resume_context
from app.ai.resume_context_attribution import classify_extracted_resume_context
from app.gmail_client import _decode_body
from app.parsing.document_extraction import (
    clean_html_text,
    elements_to_markdown,
    extract_document_text,
    prepare_gmail_parse_body,
    strip_gmail_boilerplate,
)
from app.services.role_manifest_service import RoleManifestService


def _write_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Experience", level=1)
    document.add_paragraph("Built Python and FastAPI services", style="List Bullet")
    document.add_heading("Skills", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Years"
    table.cell(1, 0).text = "PostgreSQL"
    table.cell(1, 1).text = "5"
    document.save(path)


def _convert_with_libreoffice(source: Path, output_format: str) -> Path:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", output_format, "--outdir", str(source.parent), str(source)],
        check=True,
        capture_output=True,
        text=True,
        timeout=60,
    )
    return source.with_suffix(f".{output_format}")


def _encoded(value: str) -> str:
    return base64.urlsafe_b64encode(value.encode()).decode()


class DocumentExtractionTests(unittest.TestCase):
    def test_elements_to_markdown_preserves_structure_and_drops_boilerplate(self) -> None:
        markdown = elements_to_markdown(
            [
                Header("Confidential"),
                Title("Experience", metadata=ElementMetadata(category_depth=0)),
                Title("Backend", metadata=ElementMetadata(category_depth=1)),
                ListItem("Built APIs"),
                NarrativeText("Owned production services."),
                Table(
                    "Skill Years Python 5",
                    metadata=ElementMetadata(
                        text_as_html=(
                            "<table><tr><th>Skill</th><th>Years</th></tr>"
                            "<tr><td>Python</td><td>5</td></tr></table>"
                        )
                    ),
                ),
                Footer("Unsubscribe"),
            ]
        )

        self.assertIn("## Experience", markdown)
        self.assertIn("### Backend", markdown)
        self.assertIn("- Built APIs", markdown)
        self.assertIn("| Skill | Years |", markdown)
        self.assertNotIn("Confidential", markdown)
        self.assertNotIn("Unsubscribe", markdown)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice is required to generate the PDF fixture")
    def test_extracts_real_text_pdf(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "resume.docx"
            _write_docx(source)
            path = _convert_with_libreoffice(source, "pdf")
            result = extract_document_text(str(path), path.name)

        self.assertIn("Python", result.markdown_text)
        self.assertNotIn("text extraction is limited", result.markdown_text.lower())
        self.assertTrue(result.elements)

    def test_extracts_real_docx_and_preserves_table(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            _write_docx(path)
            result = extract_document_text(str(path), path.name)

        self.assertIn("## Experience", result.markdown_text)
        self.assertIn("- Built Python and FastAPI services", result.markdown_text)
        self.assertIn("| Skill | Years |", result.markdown_text)
        self.assertIn("PostgreSQL", result.markdown_text)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice is required for legacy .doc extraction")
    def test_extracts_real_legacy_doc_through_libreoffice(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "resume.docx"
            _write_docx(source)
            path = _convert_with_libreoffice(source, "doc")
            result = extract_document_text(str(path), path.name)

        self.assertIn("Python", result.markdown_text)
        self.assertNotIn("text extraction is limited", result.markdown_text.lower())

    def test_corrupt_and_empty_documents_fall_back_without_raising(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            for suffix in (".pdf", ".docx", ".doc"):
                with self.subTest(suffix=suffix):
                    path = Path(tmp) / f"corrupt{suffix}"
                    path.write_bytes(b"not a document")
                    result = extract_document_text(str(path), path.name)
                    self.assertTrue(result.markdown_text)
            empty = Path(tmp) / "empty.pdf"
            empty.write_bytes(b"")
            result = extract_document_text(str(empty), empty.name)
            self.assertIn("text extraction is limited", result.markdown_text.lower())

    def test_section_aware_limit_marks_result_truncated(self) -> None:
        elements = [
            Title("Experience"),
            NarrativeText("Built reliable APIs. " * 20),
            Title("Education"),
            NarrativeText("Computer Science"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.pdf"
            path.write_bytes(b"placeholder")
            with patch("unstructured.partition.auto.partition", return_value=elements):
                result = extract_document_text(str(path), path.name, max_chars=140)

        self.assertTrue(result.truncated)
        self.assertLessEqual(len(result.markdown_text), 140)
        self.assertIn("Experience", result.markdown_text)
        self.assertNotIn("Education", result.markdown_text)

    def test_resume_wrapper_and_classifier_accept_markdown(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            _write_docx(path)
            markdown = extract_resume_context(str(path), path.name)

        self.assertEqual(classify_extracted_resume_context(markdown).status, "injected")
        self.assertEqual(
            classify_extracted_resume_context("Experience Python FastAPI PostgreSQL").status,
            classify_extracted_resume_context(markdown).status,
        )

    def test_clean_html_text_returns_markdown_without_css(self) -> None:
        cleaned = clean_html_text(
            """
            <html><head><style>.role { color: red; }</style></head><body>
            <h1>Java Engineer</h1><h2>Requirements</h2>
            <ul><li>Java</li><li>Spring Boot</li></ul>
            <footer>Unsubscribe</footer></body></html>
            """
        )

        self.assertIn("## Java Engineer", cleaned)
        self.assertIn("Java", cleaned)
        self.assertIn("Spring Boot", cleaned)
        self.assertNotIn("color: red", cleaned)
        self.assertNotIn("<h1>", cleaned)

    def test_single_part_html_gmail_body_uses_shared_cleaner(self) -> None:
        html = "<html><head><style>.x{display:none}</style></head><body><h1>Data Engineer</h1></body></html>"
        cleaned = _decode_body({"mimeType": "text/html", "body": {"data": _encoded(html)}})

        self.assertIn("Data Engineer", cleaned)
        self.assertNotIn("display:none", cleaned)
        self.assertNotIn("<html>", cleaned)

    def test_single_part_plain_gmail_body_is_unchanged(self) -> None:
        body = "Need a Python engineer."
        self.assertEqual(_decode_body({"mimeType": "text/plain", "body": {"data": _encoded(body)}}), body)

    def test_multipart_gmail_prefers_structured_html_and_falls_back_to_plain(self) -> None:
        payload = {
            "parts": [
                {"mimeType": "text/plain", "body": {"data": _encoded("Flat Python role")}},
                {
                    "mimeType": "text/html",
                    "body": {"data": _encoded("<html><body><h1>Python Engineer</h1><ul><li>FastAPI</li></ul></body></html>")},
                },
            ]
        }
        structured = _decode_body(payload)

        self.assertIn("Python Engineer", structured)
        self.assertIn("- FastAPI", structured)
        self.assertEqual(
            _decode_body(
                {
                    "parts": [
                        {"mimeType": "text/plain", "body": {"data": _encoded("Plain fallback")}},
                        {"mimeType": "text/html", "body": {"data": _encoded("<style>.hidden{display:none}</style>")}},
                    ]
                }
            ),
            "Plain fallback",
        )

    def test_strip_gmail_boilerplate_handles_wrapped_quote_and_preserves_reply_content(self) -> None:
        body = "\n".join(
            [
                "Rate is $70/hr and must remain.",
                "",
                "Thanks & Regards,",
                "Jamie",
                "Bench Sales Recruiter",
                "",
                "On Tue, Aug 4, 2026 at 4:20 PM Shubham <shubham@example.com>",
                "wrote:",
                "",
                "> Job Title: Lead Software Engineer",
                "> Position Summary",
                *[
                    line
                    for index in range(110)
                    for line in (f"> Requirement {index}", ">", ">")
                ],
                "> Warm Regards!",
                "> Shubham Arora",
                "> --",
                "> You received this message because you are subscribed to the Google Groups",
                "> Only C2C group.",
                "> To unsubscribe, send an email.",
                ">",
                "--",
                "You received this message because you are subscribed to the Google Groups",
                "Only C2C group.",
            ]
        )

        cleaned = strip_gmail_boilerplate(body)

        self.assertGreater(len(body.splitlines()), 250)
        self.assertLess(len(cleaned.splitlines()), 250)
        self.assertIn("Rate is $70/hr and must remain.", cleaned)
        self.assertIn("Job Title: Lead Software Engineer", cleaned)
        self.assertIn("Requirement 109", cleaned)
        self.assertNotIn("Thanks & Regards", cleaned)
        self.assertNotIn("Shubham Arora", cleaned)
        self.assertNotIn("Google Groups", cleaned)
        self.assertNotIn("\n>", cleaned)
        self.assertNotIn("wrote:", cleaned)

    def test_prepare_gmail_parse_body_composes_html_cleaning_and_dequoting(self) -> None:
        html = (
            "<html><body><p>On Tue, Aug 4, 2026 at 4:20 PM Recruiter &lt;r@example.com&gt; wrote:</p>"
            "<blockquote><h1>Data Engineer</h1><p>Python and Spark</p></blockquote></body></html>"
        )

        cleaned = prepare_gmail_parse_body(html)

        self.assertIn("Data Engineer", cleaned)
        self.assertIn("Python and Spark", cleaned)
        self.assertNotIn("<blockquote>", cleaned)
        self.assertNotIn("wrote:", cleaned)

    def test_role_manifest_line_numbers_remain_valid_for_markdown(self) -> None:
        cleaned = clean_html_text(
            "<html><body><h1>Java Engineer</h1><h2>Requirements</h2>"
            "<ul><li>Java</li><li>Spring Boot</li></ul></body></html>"
        )
        lines = cleaned.splitlines()
        captured: dict[str, str] = {}

        def provider(_system_prompt: str, user_prompt: str) -> dict[str, object]:
            captured["prompt"] = user_prompt
            return {
                "classification": "single",
                "role_count": 1,
                "confidence": 0.99,
                "shared_constraints": [],
                "roles": [
                    {
                        "index": 1,
                        "title_hint": "Java Engineer",
                        "requisition_id": "",
                        "start_line": 1,
                        "end_line": len(lines),
                        "confidence": 0.99,
                    }
                ],
            }

        result = RoleManifestService(provider=provider).detect(cleaned)

        self.assertEqual(result.status, "single")
        self.assertIn("1: ", captured["prompt"])
        self.assertIn(f"{len(lines)}: ", captured["prompt"])
        self.assertEqual(result.requirements[0].source_text, cleaned)


if __name__ == "__main__":
    unittest.main()
